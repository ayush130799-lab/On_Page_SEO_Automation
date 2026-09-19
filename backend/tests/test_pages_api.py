"""GET /api/websites/{id}/pages — the priority table, including the ``rule_id`` filter that
powers the "Most common issues" → affected-pages drilldown, and its Excel export.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.models import MemberRole, Page, SEOIssue, Severity, Website, WebsiteMember
from app.utils.url_utils import url_hash


@pytest.fixture
def site(db, member_user):
    website = Website(
        name="Acme", url="https://acme.test", domain="acme.test",
        created_by_id=member_user.id,
    )
    db.add(website)
    db.flush()
    db.add(WebsiteMember(website_id=website.id, user_id=member_user.id, role=MemberRole.OWNER))
    db.commit()
    return website


def add_page(db, website, path):
    url = f"https://acme.test{path}"
    page = Page(
        website_id=website.id, url=url, url_hash=url_hash(url), path=path,
        is_active=True, title=f"Page {path}",
    )
    db.add(page)
    db.flush()
    return page


def add_issue(db, page, rule_id, *, severity=Severity.HIGH, is_resolved=False):
    # No parent audit row needed for this filter's purposes — same convention as
    # test_impact_engine_db.py's add_issue.
    issue = SEOIssue(
        seo_audit_id=0, page_id=page.id, rule_id=rule_id, check_type=rule_id,
        severity=severity, title=rule_id.replace("_", " ").title(),
        description=f"{rule_id} is wrong", is_resolved=is_resolved,
    )
    db.add(issue)
    db.flush()
    return issue


def test_rule_id_filter_returns_only_affected_pages(db, client, member_headers, site):
    page_a = add_page(db, site, "/a")
    page_b = add_page(db, site, "/b")
    page_c = add_page(db, site, "/c")
    add_issue(db, page_a, "missing_meta_description")
    add_issue(db, page_b, "missing_meta_description")
    add_issue(db, page_c, "h1_missing")
    db.commit()

    response = client.get(
        f"/api/websites/{site.id}/pages",
        params={"rule_id": "missing_meta_description"},
        headers=member_headers,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 2
    returned_ids = {item["id"] for item in body["items"]}
    assert returned_ids == {page_a.id, page_b.id}


def test_rule_id_filter_excludes_resolved_issues(db, client, member_headers, site):
    page_a = add_page(db, site, "/a")
    add_issue(db, page_a, "missing_meta_description", is_resolved=True)
    db.commit()

    response = client.get(
        f"/api/websites/{site.id}/pages",
        params={"rule_id": "missing_meta_description"},
        headers=member_headers,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 0
    assert body["items"] == []


def test_rule_id_filter_unknown_rule_returns_empty_not_error(db, client, member_headers, site):
    add_page(db, site, "/a")
    db.commit()

    response = client.get(
        f"/api/websites/{site.id}/pages",
        params={"rule_id": "does_not_exist"},
        headers=member_headers,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 0
    assert body["items"] == []


def test_rule_id_filter_combines_with_existing_filters(db, client, member_headers, site):
    """rule_id must compose with the table's other filters, not replace them."""
    page_a = add_page(db, site, "/a")
    page_b = add_page(db, site, "/b")
    page_a.highest_severity = Severity.CRITICAL
    page_b.highest_severity = Severity.LOW
    add_issue(db, page_a, "missing_meta_description")
    add_issue(db, page_b, "missing_meta_description")
    db.commit()

    response = client.get(
        f"/api/websites/{site.id}/pages",
        params={"rule_id": "missing_meta_description", "severity": "CRITICAL"},
        headers=member_headers,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 1
    assert body["items"][0]["id"] == page_a.id


def test_rule_id_count_matches_dashboard_top_issues_page_count(db, client, member_headers, site):
    """The count shown by the "Most common issues" card must equal what the drilldown returns."""
    page_a = add_page(db, site, "/a")
    page_b = add_page(db, site, "/b")
    page_c = add_page(db, site, "/c")
    add_issue(db, page_a, "missing_meta_description")
    add_issue(db, page_b, "missing_meta_description")
    add_issue(db, page_c, "missing_meta_description", is_resolved=True)  # must not be counted
    db.commit()

    overview = client.get(f"/api/dashboard/websites/{site.id}", headers=member_headers)
    assert overview.status_code == 200
    top_issues = {i["rule_id"]: i["page_count"] for i in overview.json()["top_issues"]}
    assert top_issues["missing_meta_description"] == 2

    drilldown = client.get(
        f"/api/websites/{site.id}/pages",
        params={"rule_id": "missing_meta_description"},
        headers=member_headers,
    )
    assert drilldown.json()["total"] == top_issues["missing_meta_description"]


# ── POST /api/websites/{id}/issues/{rule_id}/export/excel ──────────────────

EXPECTED_HEADERS = [
    "URL", "Priority", "SEO", "Traffic", "Leads", "Severity", "Users", "Clicks",
    "Impressions", "CTR", "Conversions", "Issues", "Major Issues", "Intent", "AI",
]


def _load_sheet(response):
    wb = load_workbook(io.BytesIO(response.content))
    return wb.active


def test_export_issue_excel_has_exactly_the_dashboard_columns(db, client, member_headers, site):
    """Only the "Affected pages" table's columns — nothing from a URL's own detail page."""
    page_a = add_page(db, site, "/a")
    add_issue(db, page_a, "missing_meta_description")
    db.commit()

    response = client.post(
        f"/api/websites/{site.id}/issues/missing_meta_description/export/excel",
        headers=member_headers,
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    ws = _load_sheet(response)
    header_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == EXPECTED_HEADERS


def test_export_issue_excel_includes_every_affected_page_not_just_a_page(
    db, client, member_headers, site,
):
    """No pagination on the export: 5 affected pages in, 5 data rows out."""
    pages = [add_page(db, site, f"/p{i}") for i in range(5)]
    for page in pages:
        add_issue(db, page, "missing_meta_description")
    db.commit()

    response = client.post(
        f"/api/websites/{site.id}/issues/missing_meta_description/export/excel",
        headers=member_headers,
    )
    assert response.status_code == 200
    ws = _load_sheet(response)
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 5
    exported_urls = {row[0] for row in data_rows}
    assert exported_urls == {page.url for page in pages}


def test_export_issue_excel_respects_active_filters(db, client, member_headers, site):
    page_a = add_page(db, site, "/a")
    page_b = add_page(db, site, "/b")
    page_a.highest_severity = Severity.CRITICAL
    page_b.highest_severity = Severity.LOW
    add_issue(db, page_a, "missing_meta_description")
    add_issue(db, page_b, "missing_meta_description")
    db.commit()

    response = client.post(
        f"/api/websites/{site.id}/issues/missing_meta_description/export/excel",
        params={"severity": "CRITICAL"},
        headers=member_headers,
    )
    assert response.status_code == 200
    ws = _load_sheet(response)
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 1
    assert data_rows[0][0] == page_a.url


def test_export_issue_excel_filename_uses_issue_title(db, client, member_headers, site):
    page_a = add_page(db, site, "/a")
    add_issue(db, page_a, "missing_meta_description")  # title -> "Missing Meta Description"
    db.commit()

    response = client.post(
        f"/api/websites/{site.id}/issues/missing_meta_description/export/excel",
        headers=member_headers,
    )
    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert "SEO-Issue-Missing-Meta-Description.xlsx" in disposition


def test_export_issue_excel_filename_uses_explicit_title_when_given(db, client, member_headers, site):
    page_a = add_page(db, site, "/a")
    add_issue(db, page_a, "missing_meta_description")
    db.commit()

    response = client.post(
        f"/api/websites/{site.id}/issues/missing_meta_description/export/excel",
        params={"issue_title": "Meta description"},
        headers=member_headers,
    )
    assert response.status_code == 200
    assert "SEO-Issue-Meta-Description.xlsx" in response.headers["content-disposition"]


def test_export_issue_excel_unknown_rule_returns_empty_workbook_not_error(
    db, client, member_headers, site,
):
    add_page(db, site, "/a")
    db.commit()

    response = client.post(
        f"/api/websites/{site.id}/issues/does_not_exist/export/excel",
        headers=member_headers,
    )
    assert response.status_code == 200
    ws = _load_sheet(response)
    assert list(ws.iter_rows(min_row=2, values_only=True)) == []
