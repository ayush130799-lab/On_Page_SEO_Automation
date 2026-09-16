"""Excel export: data accuracy against the exact stored/denormalised values, authorization,
and behaviour on empty/invalid/large audits."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.models import (
    AIRecommendation,
    KeywordOpportunity,
    MemberRole,
    Page,
    PageIntentProfile,
    SEOAudit,
    SEOIssue,
    Severity,
    Website,
    WebsiteMember,
)
from app.models.crawl import CrawlRun

from .conftest import auth_headers, make_user

EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _seed_page(db, website, *, url="https://example.com/a", seo_score=82.0, priority_score=91.4):
    page = Page(
        website_id=website.id,
        url=url,
        url_hash=url,
        path="/a",
        seo_score=seo_score,
        priority_score=priority_score,
        priority_band="P1",
        traffic_potential_score=55.5,
        lead_potential_score=44.4,
        issue_count=1,
    )
    db.add(page)
    db.flush()

    run = CrawlRun(website_id=website.id, status="completed")
    db.add(run)
    db.flush()

    audit = SEOAudit(crawl_run_id=run.id, page_id=page.id, seo_score=seo_score)
    db.add(audit)
    db.flush()

    issue = SEOIssue(
        seo_audit_id=audit.id,
        page_id=page.id,
        rule_id="title_missing",
        check_type="title",
        category="metadata",
        severity=Severity.HIGH,
        title="Missing title tag",
        description="The page has no <title> element.",
        recommendation="Add a unique, descriptive <title>.",
        evidence={"title_length": 0},
    )
    db.add(issue)

    rec = AIRecommendation(
        website_id=website.id,
        page_id=page.id,
        provider="test",
        model="test-model",
        status="completed",
        priority="high",
        search_intent="informational",
        expected_impact="Higher CTR from search results.",
        content_quality_score=71.5,
        payload={
            "findings": [
                {
                    "issue": "Thin content",
                    "explanation": "The page has very little body text.",
                    "recommended_fix": "Expand the article with original detail.",
                    "expected_impact": "Improved topical relevance.",
                    "priority": "high",
                }
            ]
        },
    )
    db.add(rec)

    profile = PageIntentProfile(
        page_id=page.id,
        website_id=website.id,
        detected_intent="informational",
        primary_keywords=["example keyword"],
        secondary_keywords=["another keyword"],
        keyword_opportunity_score=60.0,
    )
    db.add(profile)
    db.flush()

    db.add(
        KeywordOpportunity(
            intent_profile_id=profile.id,
            page_id=page.id,
            website_id=website.id,
            keyword="example keyword",
            keyword_tier="primary",
            current_position=8.3,
        )
    )
    db.commit()
    return page


def _load(response) -> "load_workbook":
    return load_workbook(io.BytesIO(response.content))


def test_export_returns_xlsx_with_all_five_sheets(client, db, website, member_user):
    _seed_page(db, website)
    response = client.post(
        f"/api/websites/{website.id}/export/excel", headers=auth_headers(member_user)
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == EXCEL_MEDIA_TYPE
    assert "attachment" in response.headers["content-disposition"]
    assert f"SEO_Audit_Report_{website.domain}_" in response.headers["content-disposition"]

    wb = _load(response)
    assert wb.sheetnames == [
        "SEO Summary", "Page Overview", "SEO Issues", "AI Recommendations", "Keywords",
    ]


def test_export_data_matches_stored_values_exactly(client, db, website, member_user):
    page = _seed_page(db, website)
    # The dashboard reads this denormalised column directly — the export must too.
    website.average_seo_score = 82.0
    db.commit()

    response = client.post(
        f"/api/websites/{website.id}/export/excel", headers=auth_headers(member_user)
    )
    wb = _load(response)

    summary = {row[0]: row[1] for row in wb["SEO Summary"].iter_rows(values_only=True)}
    assert summary["Average SEO Score"] == 82.0
    assert summary["Total Pages"] == 1
    assert summary["High Issues"] == 1

    overview = wb["Page Overview"]
    header = [c.value for c in overview[1]]
    row = dict(zip(header, [c.value for c in overview[2]]))
    assert row["URL"] == page.url
    assert row["SEO Score"] == 82.0
    assert row["Priority"] == 91.4
    assert row["Priority Band"] == "P1"
    assert row["Traffic Potential"] == 55.5
    assert row["Lead Potential"] == 44.4
    assert row["Content Optimization Score"] == 71.5  # AIRecommendation.content_quality_score

    issues = wb["SEO Issues"]
    issue_header = [c.value for c in issues[1]]
    issue_row = dict(zip(issue_header, [c.value for c in issues[2]]))
    assert issue_row["Issue"] == "Missing title tag"
    assert issue_row["Severity"] == "HIGH"
    assert issue_row["Recommendation"] == "Add a unique, descriptive <title>."

    ai = wb["AI Recommendations"]
    ai_header = [c.value for c in ai[1]]
    ai_row = dict(zip(ai_header, [c.value for c in ai[2]]))
    assert ai_row["Issue"] == "Thin content"
    assert ai_row["AI Recommendation"] == "The page has very little body text."
    assert ai_row["Recommended Action"] == "Expand the article with original detail."
    assert ai_row["Search Intent"] == "informational"

    kw = wb["Keywords"]
    kw_header = [c.value for c in kw[1]]
    kw_row = dict(zip(kw_header, [c.value for c in kw[2]]))
    assert kw_row["Primary Keyword"] == "example keyword"
    assert kw_row["Secondary Keywords"] == "another keyword"
    assert kw_row["Current Ranking"] == 8.3


def test_export_on_website_with_no_pages_still_returns_valid_workbook(client, website, member_user):
    response = client.post(
        f"/api/websites/{website.id}/export/excel", headers=auth_headers(member_user)
    )
    assert response.status_code == 200
    wb = _load(response)
    summary = {row[0]: row[1] for row in wb["SEO Summary"].iter_rows(values_only=True)}
    assert summary["Total Pages"] == 0
    assert summary["Average SEO Score"] is None
    # Header row only — no data rows — on every tabular sheet.
    assert wb["Page Overview"].max_row == 1
    assert wb["SEO Issues"].max_row == 1


def test_export_nonexistent_website_is_not_found(client, member_user):
    response = client.post("/api/websites/999999/export/excel", headers=auth_headers(member_user))
    assert response.status_code == 404


def test_export_rejects_user_without_website_access(client, db, website):
    outsider = make_user(db, email="outsider@example.com")
    response = client.post(
        f"/api/websites/{website.id}/export/excel", headers=auth_headers(outsider)
    )
    # Deliberately 404, not 403 — see core/deps.get_website_for_read.
    assert response.status_code == 404


def test_export_requires_authentication(client, website):
    response = client.post(f"/api/websites/{website.id}/export/excel")
    assert response.status_code == 401


def test_export_allows_admin_regardless_of_membership(client, db, website, admin_user):
    _seed_page(db, website)
    response = client.post(
        f"/api/websites/{website.id}/export/excel", headers=auth_headers(admin_user)
    )
    assert response.status_code == 200


def test_export_many_pages_completes_and_row_counts_match(client, db, website, member_user):
    """Not a scale test in itself (that was verified manually against a 1,820-page/9,896-issue
    real dataset) — this proves the bulk chunked fetch joins correctly once there is more than a
    handful of rows to aggregate across pages."""
    for i in range(40):
        _seed_page(db, website, url=f"https://example.com/page-{i}", seo_score=float(60 + i))

    response = client.post(
        f"/api/websites/{website.id}/export/excel", headers=auth_headers(member_user)
    )
    assert response.status_code == 200
    wb = _load(response)
    assert wb["Page Overview"].max_row == 41  # header + 40 pages
    assert wb["SEO Issues"].max_row == 41  # one issue seeded per page
    summary = {row[0]: row[1] for row in wb["SEO Summary"].iter_rows(values_only=True)}
    assert summary["Total Pages"] == 40
    assert summary["High Issues"] == 40
